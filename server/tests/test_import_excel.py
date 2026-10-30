"""Vergelijkt de import met de door Excel berekende waarden. Draait alleen als UREN_XLSX is gezet."""
import os

import openpyxl
import pytest

from app.calc import build_year
from app.importer import import_into, parse_workbook

XLSX = os.environ.get("UREN_XLSX")
pytestmark = pytest.mark.skipif(not XLSX, reason="zet UREN_XLSX=pad/naar/salaris.xlsx")


def test_import_komt_overeen_met_excel(conn):
    data = open(XLSX, "rb").read()
    import_into(conn, parse_workbook(data))
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for year, ot_col in ((2025, 17), (2026, 17)):
        ws = wb[str(year)]
        Y = build_year(conn, year)
        for i, m in enumerate(Y["months"]):
            assert m.overuren == pytest.approx(ws.cell(11 + i, ot_col).value or 0), (year, m.naam)
        assert Y["totals"]["uren"] == pytest.approx(ws["I63"].value)
    Y = build_year(conn, 2026)
    assert Y["totals"]["vakantie_gebruikt"] == wb["2026"]["H3"].value
    assert Y["totals"]["storing_dagen"] == 92
