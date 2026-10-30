"""Export naar Excel: één tabblad per jaar, in de opzet van het oude bestand (waarden, geen formules)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calc import DAGEN, build_year

EURO = '"€" #,##0.00'
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")


def years_with_data(conn) -> list[int]:
    ys = {int(r[0][:4]) for r in conn.execute("SELECT date FROM days")}
    ys |= {r[0] for r in conn.execute("SELECT year FROM months")}
    ys |= {r[0] for r in conn.execute("SELECT year FROM years")}
    return sorted(ys)


def export_workbook(conn, years: list[int] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for year in years or years_with_data(conn):
        _sheet(wb, build_year(conn, year))
    _days_sheet(wb, conn)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(wb, Y):
    ws = wb.create_sheet(str(Y["year"]))
    s, t = Y["settings"], Y["totals"]
    info = [("Jaar", Y["year"]), ("Contract uren p/w", s["contract_hours"]),
            ("Overuren schaal 1", s["rate1"]), ("Overuren schaal 2", s["rate2"]),
            ("Storing per dag", s["storing_rate"]), ("Vakantiedagen totaal", s["vakantie_totaal"]),
            ("Vakantiedagen gebruikt", t["vakantie_gebruikt"]), ("Vakantiedagen over", t["vakantie_over"]),
            ("Totaal uren", t["uren"])]
    for i, (k, v) in enumerate(info, start=1):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, v)
    for r in (3, 4, 5):
        ws.cell(r, 2).number_format = EURO

    hr = len(info) + 2
    heads = ["Week"] + DAGEN + ["Totaal", "Overuren", "Maand", "Dagen storing", "Vakantiedagen"]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(hr, c, h)
        cell.font, cell.fill = BOLD, HEAD
    r = hr
    for w in Y["weeks"]:
        r += 1
        ws.cell(r, 1, w.week)
        for i, h in enumerate(w.hours):
            if h is not None:
                ws.cell(r, 2 + i, h)
        ws.cell(r, 9, w.total or None)
        ws.cell(r, 10, w.overuren)
        ws.cell(r, 11, Y["months"][w.month - 1].naam)
        ws.cell(r, 12, w.storing_days or None)
        ws.cell(r, 13, w.vakantie_days or None)

    mc = 15
    mheads = ["Maand", "Salaris", "Nettoloon", "Overuur schaal", "Uren", "Overuren", "Gewerkt (zonder factor)", "Netto € overuren",
              "Storingsdagen", "Storingsdienst €", "Totaal netto"]
    for c, h in enumerate(mheads):
        cell = ws.cell(hr, mc + c, h)
        cell.font, cell.fill = BOLD, HEAD
    for i, m in enumerate(Y["months"], start=1):
        vals = [m.naam, m.bruto, m.netto, m.schaal, m.uren, m.overuren, m.gewerkt, m.overuren_euro,
                m.storing_days, m.storing_euro, m.totaal_netto]
        for c, v in enumerate(vals):
            cell = ws.cell(hr + i, mc + c, v)
            if c in (1, 2, 7, 9, 10):
                cell.number_format = EURO
    tr = hr + 13
    ws.cell(tr, mc, "Totaal").font = BOLD
    for c, v in ((1, t["bruto"]), (2, t["netto"]), (4, t["uren"]), (5, t["overuren"]),
                 (6, sum(m.gewerkt for m in Y["months"])), (7, t["overuren_euro"]),
                 (8, t["storing_dagen"]), (9, t["storing_euro"]), (10, t["totaal_netto"])):
        cell = ws.cell(tr, mc + c, v)
        cell.font = BOLD
        if c in (1, 2, 7, 9, 10):
            cell.number_format = EURO

    if Y["payouts"]:
        pr = tr + 3
        ws.cell(pr, mc, "Gehad overuren").font = BOLD
        for j, p in enumerate(Y["payouts"], start=1):
            ws.cell(pr + j, mc, p["date"])
            ws.cell(pr + j, mc + 1, p["amount"]).number_format = EURO

    for c in range(1, mc + len(mheads)):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.column_dimensions["A"].width = 22
    for c in range(1, len(heads) + 1):
        ws.cell(hr, c).alignment = Alignment(wrap_text=True)
    ws.freeze_panes = ws.cell(hr + 1, 2)


def _days_sheet(wb, conn):
    ws = wb.create_sheet("Dagen")
    heads = ["Datum", "Uren", "Storing", "Vakantie", "Telt x (zon-/feestdag)", "Notitie"]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = BOLD, HEAD
    for r, d in enumerate(conn.execute("SELECT * FROM days ORDER BY date"), start=2):
        ws.cell(r, 1, d["date"])
        ws.cell(r, 2, d["hours"])
        ws.cell(r, 3, "ja" if d["storing"] else None)
        ws.cell(r, 4, d["vakantie"] or None)
        ws.cell(r, 5, d["factor"] if d["factor"] != 1 else None)
        ws.cell(r, 6, d["note"])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["F"].width = 40
