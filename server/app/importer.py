"""Import van het oude Excel-bestand (tabbladen per jaar, bv. '2023' t/m '2026').

Kolommen worden op koptekst gezocht, zodat de iets verschillende indeling per jaar werkt.
De maand per week wordt overgenomen zoals Excel die berekende, zodat maandtotalen gelijk blijven.
Storings- en vakantiedagen staan in het Excel per week; die worden als week-aantal opgeslagen.
"""
from __future__ import annotations

import datetime as dt
import io

import openpyxl

from .calc import DAGEN, MAANDEN, default_month, weeks_in_year

_SETTING_LABELS = {
    "contract uren p/w": "contract_hours",
    "betaling overuren": "rate1",
}
_G_LABELS = {
    "bedrag per dag storing": "storing_rate",
    "dagen totaal": "vakantie_totaal",
}


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def parse_workbook(data: bytes) -> dict:
    wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    years = {}
    for name in wb_v.sheetnames:
        if not name.strip().isdigit():
            continue
        years[int(name)] = _parse_sheet(int(name), wb_v[name], wb_f[name])
    return years


def _parse_sheet(year: int, ws, ws_f) -> dict:
    out = dict(year=year, settings={}, days=[], weeks=[], months=[], payouts=[])
    # instellingen (kolom A label, waarde in C of B)
    for r in range(1, 10):
        label = _norm(ws.cell(r, 1).value)
        for key, field in _SETTING_LABELS.items():
            if label.startswith(key):
                val = _num(ws.cell(r, 3).value)
                if val is None:
                    val = _num(ws.cell(r, 2).value)
                if val is not None:
                    out["settings"][field] = float(val)
                if field == "rate1":
                    rate2 = _num(ws.cell(r + 1, 3).value)
                    if rate2 is not None:
                        out["settings"]["rate2"] = float(rate2)
        if label.startswith("loonheffingskorting"):
            out["settings"]["loonheffingskorting"] = 0 if _norm(ws.cell(r, 2).value) == "nee" else 1
        if label.startswith("situatie"):
            out["settings"]["situatie"] = "standard" if _norm(ws.cell(r, 2).value).startswith("jonger") else "aow"
        for c in range(4, 12):
            g = _norm(ws.cell(r, c).value)
            for key, field in _G_LABELS.items():
                if g.startswith(key):
                    val = _num(ws.cell(r, c + 1).value)
                    if val is not None:
                        out["settings"][field] = float(val)

    # kopregel van de weektabel
    header_row = next(r for r in range(1, 20) if _norm(ws.cell(r, 1).value) == "week")
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = _norm(ws.cell(header_row, c).value)
        if h and h not in cols:
            cols[h] = c
        elif h == "maand":
            cols["maand2"] = c
    day_cols = [cols[d.lower()] for d in DAGEN]
    nweeks = weeks_in_year(year)

    for r in range(header_row + 1, header_row + 60):
        w = ws.cell(r, 1).value
        if not isinstance(w, int) or not 1 <= w <= nweeks:
            continue
        has_hours = False
        for i, c in enumerate(day_cols):
            h = _num(ws.cell(r, c).value)
            if h is not None:
                has_hours = True
                out["days"].append(dict(date=dt.date.fromisocalendar(year, w, i + 1).isoformat(), hours=float(h)))
        storing = _num(ws.cell(r, cols["dagen storing"]).value) if "dagen storing" in cols else None
        vak = _num(ws.cell(r, cols["vakantiedagen"]).value) if "vakantiedagen" in cols else None
        month = None
        k = _norm(ws.cell(r, cols["maand"]).value) if "maand" in cols else ""
        if k in MAANDEN and (has_hours or storing or vak):
            m = MAANDEN.index(k) + 1
            if m != default_month(year, w):
                month = m
        if month or storing or vak:
            out["weeks"].append(dict(week=w, month=month, storing_extra=int(storing or 0),
                                     vakantie_extra=float(vak or 0)))

    # maandtabel
    mc = cols.get("maand2")
    if mc:
        hdr = {_norm(ws.cell(header_row, c).value): c for c in range(mc, mc + 10)}
        for r in range(header_row + 1, header_row + 14):
            naam = _norm(ws.cell(r, mc).value)
            if naam not in MAANDEN:
                continue
            row = dict(month=MAANDEN.index(naam) + 1)
            if "salaris" in hdr:
                row["bruto"] = _num(ws.cell(r, hdr["salaris"]).value)
            if "nettoloon" in hdr:
                row["netto"] = _num(ws.cell(r, hdr["nettoloon"]).value)
            row["schaal"] = int(_num(ws.cell(r, hdr["overuur schaal"]).value) or 1) if "overuur schaal" in hdr else 1
            out["months"].append(row)

    # gehad overuren (datum + bedrag)
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(r, c).value) == "gehad overuren":
                rr = r + 2
                while isinstance(ws.cell(rr, c).value, (dt.datetime, dt.date)):
                    amount = _num(ws.cell(rr, c + 1).value)
                    if amount is not None:
                        d = ws.cell(rr, c).value
                        out["payouts"].append(dict(date=(d.date() if isinstance(d, dt.datetime) else d).isoformat(),
                                                   amount=float(amount)))
                    rr += 1
    return out


def import_into(conn, parsed: dict, years: list[int] | None = None, now: str | None = None) -> dict:
    now = now or dt.datetime.now().isoformat(timespec="seconds")
    stats = {}
    for year, data in sorted(parsed.items()):
        if years and year not in years:
            continue
        s = data["settings"]
        if s:
            cur = conn.execute("SELECT * FROM years WHERE year=?", (year,)).fetchone()
            base = dict(cur) if cur else {}
            base.update(s)
            base["year"] = year
            keys = [k for k in ("year", "contract_hours", "rate1", "rate2", "storing_rate", "vakantie_totaal",
                                "loonheffingskorting", "situatie") if k in base]
            conn.execute(f"INSERT OR REPLACE INTO years({','.join(keys)}) VALUES({','.join('?' * len(keys))})",
                         [base[k] for k in keys])
        for d in data["days"]:
            conn.execute(
                "INSERT INTO days(date, hours, updated_at) VALUES(?,?,?) "
                "ON CONFLICT(date) DO UPDATE SET hours=excluded.hours, updated_at=excluded.updated_at",
                (d["date"], d["hours"], now))
        for w in data["weeks"]:
            conn.execute("INSERT OR REPLACE INTO weeks(iso_year, week, month, storing_extra, vakantie_extra) "
                         "VALUES(?,?,?,?,?)", (year, w["week"], w["month"], w["storing_extra"], w["vakantie_extra"]))
        for m in data["months"]:
            conn.execute("INSERT OR REPLACE INTO months(year, month, bruto, netto, schaal) VALUES(?,?,?,?,?)",
                         (year, m["month"], m.get("bruto"), m.get("netto"), m.get("schaal", 1)))
        conn.execute("DELETE FROM payouts WHERE substr(date,1,4)=?", (str(year),))
        for p in data["payouts"]:
            conn.execute("INSERT INTO payouts(date, amount, note) VALUES(?,?,?)",
                         (p["date"], p["amount"], "import Excel"))
        stats[year] = dict(dagen=len(data["days"]), weken=len(data["weeks"]), maanden=len(data["months"]),
                           uitbetalingen=len(data["payouts"]))
    return stats
